"""
content.py — turn source material into template-ready HTML.

Lifted from the desktop db_creator.py so the web app parses Word docs and typed
headers exactly the way the desktop tool does. Nothing here touches Selenium,
Tkinter, or the network: give it a .docx path or a list of (header, content)
pairs and it hands back (title, body_html, sections).

    title       "Week 1 Discussion: Medicare Matters"
    body_html   everything in order, for single-box templates
    sections    prompt / objective / instructions plus any section the source
                names itself, each ready to drop into the matching region
"""

import re as _re
import html as _htmlesc
from pathlib import Path

# Standalone labels in the Word doc that should NOT appear in the Canvas box
_SKIP_LABELS = {
    "scenario", "prompt", "discussion prompt", "overview", "background",
    "instructions", "directions", "case study", "case scenario",
}


# Section labels the Word doc may use to name each block of content. These map
# onto the matching regions of the Canvas template.
_OVERVIEW_LABELS = {
    "discussion prompt", "prompt", "scenario", "case study", "case scenario",
    "discussion overview", "overview", "background", "summary", "description",
}
_INSTRUCTION_LABELS = {
    "instructions", "discussion instructions", "instructions for students",
    "directions", "task", "your task", "requirements", "questions",
    "discussion questions", "steps",
}
_OBJECTIVE_LABELS = {"objective", "objectives", "assignment objective", "purpose"}
_ALL_SECTION_LABELS = (
    _OVERVIEW_LABELS | _INSTRUCTION_LABELS | _OBJECTIVE_LABELS
)

# Matches "[Assignment Title Here]", "[insert objective here]", etc.
_BRACKET_RE = _re.compile(r"\[[^\[\]]{0,240}\]")


def _bracket_placeholders(html: str):
    """Every [ ... ] placeholder left in a chunk of HTML, text only."""
    text = _re.sub(r"<[^>]+>", " ", html or "")
    text = text.replace("&nbsp;", " ").replace("\u00a0", " ")
    return [m.group(0) for m in _BRACKET_RE.finditer(text)]


def _run_to_html(run):
    """Convert one Word run to inline HTML, preserving bold/italic/underline."""
    txt = _htmlesc.escape(run.text or "")
    if not txt.strip():
        return txt
    if run.underline:
        txt = f"<u>{txt}</u>"
    if run.italic:
        txt = f"<em>{txt}</em>"
    if run.bold:
        txt = f"<strong>{txt}</strong>"
    return txt


def _para_to_html(p):
    """
    Convert a Word paragraph to HTML, including hyperlinks.

    python-docx leaves runs nested inside a <w:hyperlink> out of p.runs, so
    walking p.runs alone silently drops every linked phrase. This walks the
    paragraph XML in document order instead and rebuilds the <a> tags.
    """
    from docx.text.run import Run
    from docx.oxml.ns import qn as _qn

    R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    parts = []

    for child in p._p.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "r":
            parts.append(_run_to_html(Run(child, p)))
        elif tag == "hyperlink":
            inner = "".join(_run_to_html(Run(r, p)) for r in child.findall(_qn("w:r")))
            if not inner.strip():
                continue
            url = ""
            rid = child.get(R_ID)
            if rid:
                try:
                    url = p.part.rels[rid].target_ref or ""
                except Exception:
                    url = ""
            anchor = child.get(_qn("w:anchor"))
            if not url and anchor:
                url = "#" + anchor
            parts.append(f'<a href="{_htmlesc.escape(url, quote=True)}">{inner}</a>' if url else inner)

    out = "".join(parts).strip()
    return out or _htmlesc.escape((p.text or "").strip())


def _para_is_all_bold(p):
    runs = [r for r in p.runs if (r.text or "").strip()]
    return bool(runs) and all(bool(r.bold) for r in runs)


def _list_info(doc, p):
    """Return (is_list_item, is_ordered) for a Word paragraph."""
    style = (getattr(p.style, "name", "") or "").lower()
    numPr = None
    try:
        pPr = p._p.pPr
        numPr = pPr.numPr if pPr is not None else None
    except Exception:
        numPr = None

    if numPr is None:
        if "list number" in style:
            return True, True
        if "list bullet" in style:
            return True, False
        return False, False

    # Numbered paragraph — resolve the real numbering format when possible
    ordered = "number" in style
    try:
        num_id = numPr.numId.val
        ilvl = numPr.ilvl.val if numPr.ilvl is not None else 0
        numbering = doc.part.numbering_part.element
        nums = numbering.xpath(f"./w:num[@w:numId='{num_id}']")
        abstract_id = nums[0].xpath("./w:abstractNumId/@w:val")[0]
        fmt = numbering.xpath(
            f"./w:abstractNum[@w:abstractNumId='{abstract_id}']"
            f"/w:lvl[@w:ilvl='{ilvl}']/w:numFmt/@w:val"
        )
        if fmt:
            ordered = str(fmt[0]).lower() != "bullet"
    except Exception:
        pass
    return True, ordered


def _title_from_filename(path):
    stem = Path(path).stem.replace("_", " ").replace("-", " ").strip()
    return _re.sub(r"\s+", " ", stem)


def _render_items(items):
    """Turn a list of parsed paragraph dicts into HTML, grouping runs of list items."""
    parts, open_list = [], None
    for it in items:
        if it["list"]:
            tag = "ol" if it["ordered"] else "ul"
            if open_list != tag:
                if open_list:
                    parts.append(f"</{open_list}>")
                parts.append(f"<{tag}>")
                open_list = tag
            parts.append(f"<li>{it['html']}</li>")
        else:
            if open_list:
                parts.append(f"</{open_list}>")
                open_list = None
            parts.append(f"<p>{it['html']}</p>")
    if open_list:
        parts.append(f"</{open_list}>")
    return "\n".join(parts)


def _is_section_label(it):
    """
    True when a paragraph is acting as a section label rather than content.

    Three ways to mark one in the Word doc:
      Instructions        a name the tool already knows
      Format              any short standalone line, matched to the template
                          heading of the same name
      [Objective]         wrapped in brackets, matched to the template's
                          [ ... ] placeholder containing that word
    """
    if it["list"]:
        return False
    text = it["text"].strip()
    if not text or len(text.split()) > 6:
        return False
    if it.get("label_hint"):
        return True
    if text.startswith("[") and text.endswith("]"):
        return True
    label = text.rstrip(":").lower()
    if label in _ALL_SECTION_LABELS or label in _SKIP_LABELS:
        return True
    # A short bold line that isn't a sentence is a heading in every doc style
    return bool(it.get("bold")) and not text.endswith((".", "!", "?", ","))


def _is_label(it, labels):
    """True when a paragraph is just a section label like 'Assignment Overview'."""
    if it["list"]:
        return False
    label = it["text"].strip().rstrip(":").lower()
    return label in labels and len(it["text"].split()) <= 5


def parse_discussion_docx(path):
    """
    Parse a Week discussion Word doc.

    Returns (title, body_html, sections).
      title      e.g. "Week 1 Discussion: Medicare Matters"
      body_html  everything, for templates with a single prompt box
      sections   dict of 'prompt', 'objective', 'instructions' HTML plus any
                 extra sections the doc names itself, each ready to drop into
                 the matching template region
    """
    from docx import Document

    doc = Document(str(path))

    items = []
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        is_list, ordered = _list_info(doc, p)
        # A soft line break (Shift+Enter) inside one Word paragraph is how many
        # docs write "Label<break>body text". Split those into separate items so
        # the label can be recognised instead of being glued to the sentence.
        html = _para_to_html(p)
        text_parts = [t for t in text.split("\n")]
        html_parts = html.split("\n")
        if len(html_parts) != len(text_parts):
            text_parts, html_parts = [text], [html]
        segments = [(t.strip(), h.strip())
                    for t, h in zip(text_parts, html_parts) if t.strip()]
        para_bold = _para_is_all_bold(p)
        for n, (seg_text, seg_html) in enumerate(segments):
            items.append({
                "text": seg_text,
                "html": seg_html,
                "list": is_list,
                "ordered": ordered,
                "bold": para_bold,
                # First of several lines in one paragraph, short, and not a
                # sentence: that is a section label in every doc style.
                "label_hint": (n == 0 and len(segments) > 1 and not is_list
                               and len(seg_text.split()) <= 6
                               and not seg_text.rstrip().endswith((".", "!", "?", ","))),
            })

    empty_sections = {"_week": None, "_topic": "", "_regions": [],
                      "prompt": "", "objective": "", "instructions": ""}
    if not items:
        return _title_from_filename(path), "", dict(empty_sections)

    # ── Title: pull the week number and the topic out of the first few lines ──
    week_num = None
    topic = None
    header_idx = -1
    title_idx = -1

    for i, it in enumerate(items[:6]):
        if it["list"]:
            continue
        if week_num is None:
            m = _re.search(r"week\s*#?\s*(\d+)", it["text"], _re.I)
            if m:
                week_num = int(m.group(1))
                header_idx = max(header_idx, i)
        if topic is None:
            m2 = _re.match(
                r"^\s*(?:week\s*#?\s*\d+\s*)?discussion\s*[:\-\u2010-\u2015]\s*(.+)$",
                it["text"], _re.I
            )
            if m2:
                topic = m2.group(1).strip()
                header_idx = max(header_idx, i)

    if week_num is None:
        m = _re.search(r"week\s*_?\s*#?\s*(\d+)", Path(path).stem, _re.I)
        if m:
            week_num = int(m.group(1))

    # No "Discussion:" line? Take the first short standalone line as the topic.
    if topic is None:
        for i, it in enumerate(items[:8]):
            if it["list"]:
                continue
            label = it["text"].strip().rstrip(":").lower()
            if label in _ALL_SECTION_LABELS or label in _SKIP_LABELS:
                continue
            words = it["text"].split()
            if not (2 <= len(words) <= 12):
                continue
            if it["text"].rstrip().endswith((".", "!", "?", ":")):
                continue
            if _re.match(r"^\s*(in this|the objective|post|respond|you will)\b",
                         it["text"], _re.I):
                continue
            topic = it["text"].strip()
            title_idx = i
            break

    if topic and week_num:
        title = f"Week {week_num} Discussion: {topic}"
    elif topic:
        title = f"Discussion: {topic}"
    elif week_num:
        title = f"Week {week_num} Discussion"
    else:
        title = _title_from_filename(path)

    # ── Body ─────────────────────────────────────────────────────────────────
    body = items[header_idx + 1:] if header_idx >= 0 else items[:]
    if title_idx >= 0:
        offset = title_idx - (header_idx + 1 if header_idx >= 0 else 0)
        if 0 <= offset < len(body):
            body.pop(offset)

    # Docs often name the week up top and repeat the bare topic further down.
    # That repeat is a heading, not content, so it must not land in a section.
    def _same_line(a, b):
        norm = lambda t: _re.sub(r"[\s:.\-\u2010-\u2015]+$", "",
                                 " ".join((t or "").split()).lower())
        return bool(norm(a)) and norm(a) == norm(b)

    body = [it for it in body
            if it["list"] or not (_same_line(it["text"], title)
                                  or _same_line(it["text"], topic or ""))]

    # ── Group under whichever section labels the doc uses ────────────────────
    # Any short standalone line acts as a label. Write "Response Requirements"
    # in the doc and it routes to the template heading of the same name; wrap
    # it, "[Prompt]", and it routes to the template's matching [ ... ] spot.
    groups, current, bucket = [], None, []
    for it in body:
        if _is_section_label(it):
            groups.append((current, bucket))
            current, bucket = it["text"].strip(), []
            continue
        bucket.append(it)
    groups.append((current, bucket))
    groups = [(lab, its) for lab, its in groups if its]

    prompt_items, instruction_items, objective_item, loose = [], [], None, []
    extras = []
    # Where each section sat in the doc, and the header the doc gave it. The
    # merge fills whichever section the template names first, so without this
    # the rest get added after it and the finished box reads out of order with
    # its headers missing.
    order = {}
    labels = {}
    # Two sections of the doc can route to the same region, so each one is also
    # kept on its own; joining them would lose the later header and its place.
    chunks = []

    def _mark(key, idx, shown=""):
        if key not in order:
            order[key] = idx
            if shown:
                labels[key] = shown

    def _chunk(route, idx, shown, items):
        if items:
            chunks.append({"route": route, "order": idx, "label": shown or "",
                           "items": items})

    for gi, (label, its) in enumerate(groups):
        shown = (label or "").strip()
        raw = (label or "").strip().rstrip(":")
        bracketed = raw.startswith("[") and raw.endswith("]")
        key = (raw[1:-1] if bracketed else raw).strip().lower()
        if label is None:
            loose.append((gi, its))
        elif key in _OVERVIEW_LABELS:
            prompt_items.extend(its)
            _mark("prompt", gi, shown)
            _chunk("prompt", gi, shown, its)
        elif key in _INSTRUCTION_LABELS:
            instruction_items.extend(its)
            _mark("instructions", gi, shown)
            _chunk("instructions", gi, shown, its)
        elif key in _OBJECTIVE_LABELS and its:
            if objective_item is None:
                objective_item = its[0]
                _mark("objective", gi, shown)
                _chunk("objective", gi, shown, its[:1])
            else:
                instruction_items.extend(its[:1])
                _mark("instructions", gi, shown)
                _chunk("instructions", gi, shown, its[:1])
            instruction_items.extend(its[1:])
            if len(its) > 1:
                _mark("instructions", gi, shown)
                _chunk("instructions", gi, shown, its[1:])
        else:
            extras.append({"label": raw, "keyword": key, "bracket": bracketed,
                           "items": its, "order": gi, "display": shown})

    loose_order = loose[0][0] if loose else None
    loose = [it for _gi, its in loose for it in its]

    if objective_item is None:
        for pool in (loose, instruction_items, prompt_items):
            for it in list(pool):
                if not it["list"] and _re.search(
                    r"\bobjective of this (discussion|assignment|task)\b",
                    it["text"], _re.I
                ):
                    objective_item = it
                    pool.remove(it)
                    break
            if objective_item:
                break

    # Unlabelled docs: the scenario comes first, the questions follow
    if loose:
        if not prompt_items and not instruction_items:
            prompt_items = loose
            if loose_order is not None:
                order["prompt"] = loose_order
            _chunk("prompt", loose_order or 0, "", loose)
        elif not instruction_items:
            instruction_items = loose
            if loose_order is not None:
                order["instructions"] = loose_order
            _chunk("instructions", loose_order or 0, "", loose)
        else:
            instruction_items = loose + instruction_items
            if loose_order is not None:
                order["instructions"] = min(loose_order,
                                            order.get("instructions", loose_order))
            _chunk("instructions", loose_order or 0, "", loose)

    sections = {
        "_week": week_num,
        "_topic": topic or title,
        "_order": order,
        "_labels": labels,
        "_chunks": [{"route": c["route"], "order": c["order"],
                     "label": c["label"], "html": _render_items(c["items"])}
                    for c in sorted(chunks, key=lambda c: c["order"])],
        "_regions": [{"label": e["label"], "keyword": e["keyword"],
                      "bracket": e["bracket"], "order": e["order"],
                      "display": e.get("display") or e["label"],
                      "html": _render_items(e["items"])}
                     for e in extras],
        "prompt": _render_items(prompt_items),
        "objective": _render_items([objective_item]) if objective_item else "",
        "instructions": _render_items(instruction_items),
    }

    ordered = prompt_items + ([objective_item] if objective_item else []) + instruction_items
    body_html = _render_items(ordered)

    return title, body_html, sections


# ── Typed headers + pasted content ───────────────────────────────────────────
# Same idea as the Word doc parser, but you type the header yourself and paste
# the content that belongs under it. Each header routes to the template region
# of that name exactly the way a labelled Word section does.

_TITLE_LABELS = {
    "title", "discussion title", "discussion name", "topic", "discussion topic",
    "name", "week title",
}

_BULLET_RE = _re.compile(
    r"^\s*(?:([-*\u2022\u00b7\u25cf\u25aa\u25e6\u2043\u2013\u2014])"
    r"|(\d{1,2}\s*[.)])"
    r"|([a-zA-Z]\s*\)))\s+(.+)$"
)

# Word's second-level bullet copies out as a lone "o" and a tab. Requiring the
# tab keeps real words that start with o from being mistaken for markers.
_SUB_BULLET_RE = _re.compile(r"^o\t\s*(.+)$")

# A line Word indented but gave no marker: a tab, or two or more spaces
_INDENT_RE = _re.compile(r"^(?:\t| {2,}|\u00a0{2,})")


def _norm_label(text: str) -> str:
    """Lowercase, unbracketed, no trailing punctuation — for comparing headers."""
    t = (text or "").replace("\u00a0", " ")
    t = _re.sub(r"\s+", " ", t).strip()
    if t.startswith("[") and t.endswith("]"):
        t = t[1:-1].strip()
    return _re.sub(r"[\s:.\-\u2010-\u2015]+$", "", t).lower()


def _inline_text_to_html(line: str) -> str:
    """Escape a pasted line and honour **bold**, *italic*, and bare URLs."""
    out = _htmlesc.escape(line.strip())
    out = _re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    out = _re.sub(r"(?<!\*)\*(?!\s)([^*]+?)(?<!\s)\*(?!\*)", r"<em>\1</em>", out)
    out = _re.sub(r"(https?://[^\s<]+?)(?=[.,;:)\]]?(?:\s|$))",
                  r'<a href="\1">\1</a>', out)
    return out


def _looks_like_html(text: str) -> bool:
    t = (text or "").strip()
    return bool(_re.match(r"^<(p|ul|ol|div|h[1-6]|table|blockquote)\b", t, _re.I)
                and _re.search(r"</\w+>\s*$", t))


_LIST_MARK_RE = _re.compile(
    r"^\s*(?:[-*\u2022\u00b7\u25cf\u25aa\u25e6]|\d{1,2}\s*[.)])\s+")


def _question_runs(lines):
    """
    Line numbers that form a list even though nothing marks them as one.

    Word often pastes a bulleted list as bare lines with no glyph and no
    indent. Two or more question lines in a row, right after a line ending in
    a colon, is that shape and nothing else, so they are treated as bullets.
    """
    flags, n, i = set(), len(lines), 0
    while i < n:
        if lines[i].strip().endswith(":"):
            j, run = i + 1, []
            while j < n:
                t = lines[j].strip()
                if not t:
                    j += 1
                    continue
                if t.endswith("?"):
                    run.append(j)
                    j += 1
                    continue
                break
            if len(run) >= 2:
                flags.update(run)
                i = j
                continue
        i += 1
    return flags


def _text_to_html(text: str) -> str:
    """
    Turn pasted plain text into template-ready HTML.

    Blank-line or single-line breaks become paragraphs; lines that start with a
    bullet character or a number become list items, grouped into <ul>/<ol>.
    Text that is already HTML is passed straight through.
    """
    raw = (text or "").strip()
    if not raw:
        return ""
    if _looks_like_html(raw):
        return raw

    items = []
    run, after_colon = False, False
    all_lines = raw.splitlines()
    q_runs = _question_runs(all_lines)
    for _idx, line in enumerate(all_lines):
        s = line.strip()
        if not s:
            run = False
            continue
        # Indentation is read before stripping, since that is all Word leaves
        # behind when its bullet glyphs do not survive the copy.
        indented = bool(_INDENT_RE.match(line))

        m = _BULLET_RE.match(s)
        if m:
            body = m.group(4).strip()
            ordered = bool(m.group(2) or m.group(3))
            items.append({"html": _inline_text_to_html(body),
                          "list": True, "ordered": ordered})
            run, after_colon = True, False
            continue

        sub = _SUB_BULLET_RE.match(s)
        if sub:
            items.append({"html": _inline_text_to_html(sub.group(1).strip()),
                          "list": True, "ordered": False})
            run, after_colon = True, False
            continue

        # An unmarked, unindented run of questions under a "...:" line
        if _idx in q_runs:
            items.append({"html": _inline_text_to_html(s),
                          "list": True, "ordered": False})
            run, after_colon = True, False
            continue

        # An indented line right after a line ending in ":" starts a list, and
        # every indented line after it stays in that list.
        if indented and (after_colon or run):
            items.append({"html": _inline_text_to_html(s),
                          "list": True, "ordered": False})
            run, after_colon = True, False
            continue

        items.append({"html": _inline_text_to_html(s),
                      "list": False, "ordered": False})
        run = False
        after_colon = s.endswith(":")
    return _render_items(items)


# ── Banner heading clone ─────────────────────────────────────────────────────
# A header the template has no spot for used to come out as a bare <h3>, which
# lands unstyled in the middle of the page. Instead, the template's own banner
# heading block is cloned: same row, same icon wrapper, same white-on-red type,
# with the new header text swapped in.

NEW_HEADING_ICON = "find_in_page"      # Material icon used on cloned banners
NEW_HEADING_BG = "#b20000"             # background-color that marks the banner row

_BANNER_BG_RE = _re.compile(
    r"background-color:\s*" + _re.escape(NEW_HEADING_BG), _re.I)


def _div_slice(html: str, start: int) -> str:
    """The full <div>…</div> that begins at `start`, nesting included."""
    depth = 0
    for m in _re.finditer(r"<\s*(/?)div\b", html[start:], _re.I):
        if m.group(1):
            depth -= 1
            if depth == 0:
                gt = html.find(">", start + m.end())
                return html[start:gt + 1] if gt != -1 else ""
        else:
            depth += 1
    return ""


def _find_banner_wrapper(template_html: str):
    """
    Locate the template's banner heading.

    Returns (wrapper_html, row_open_tag): the icon + heading wrapper, and the
    opening tag of the row that carries the banner colour, so the clone keeps
    the template's own padding and background.
    """
    html = template_html or ""
    for m in _re.finditer(r"<div\b[^>]*>", html, _re.I):
        tag = m.group(0)
        if "loree-iframe-content-row" not in tag or not _BANNER_BG_RE.search(tag):
            continue
        row = _div_slice(html, m.start())
        if not row:
            continue
        wm = _re.search(r"<div\b[^>]*special-element-wrapper[^>]*>", row, _re.I)
        if not wm:
            continue
        wrapper = _div_slice(row, wm.start())
        if wrapper and _re.search(r"<h[1-6]\b", wrapper, _re.I):
            return wrapper, tag
    return "", ""


def _strip_leaf_lists(html: str) -> str:
    """Remove <ul>/<ol> blocks, innermost first, so nesting is handled."""
    leaf = _re.compile(r"(?is)<(ul|ol)\b[^>]*>(?:(?!<(?:ul|ol)\b).)*?</\1>")
    for _ in range(6):
        stripped = leaf.sub("", html)
        if stripped == html:
            break
        html = stripped
    return html


def banner_heading_text(template_html: str) -> str:
    """The heading the template's banner currently carries, plain text."""
    wrapper, _row = _find_banner_wrapper(template_html)
    if not wrapper:
        return ""
    m = _re.search(r"(?is)<h[1-6]\b[^>]*>(.*?)</h[1-6]>", wrapper)
    if not m:
        return ""
    txt = _re.sub(r"<[^>]+>", " ", m.group(1))
    txt = _htmlesc.unescape(txt).replace("\u00a0", " ")
    return _re.sub(r"\s+", " ", txt).strip()


def banner_heading_html(template_html: str, header_text: str,
                        icon: str = NEW_HEADING_ICON) -> str:
    """
    Build a banner-styled heading for a header the template does not have.

    Falls back to a plain <h3> when the template has no banner to copy, so this
    is always safe to call.
    """
    text = (header_text or "").strip().strip("[]").strip()
    esc = _htmlesc.escape(text)
    wrapper, row_open = _find_banner_wrapper(template_html)
    if not wrapper:
        return f"<h3>{esc}</h3>"

    block = wrapper

    # 1. New header text in the cloned heading
    block = _re.sub(r"(?is)(<h([1-6])\b[^>]*>)(.*?)(</h\2>)",
                    lambda m: m.group(1) + esc + m.group(4), block, count=1)

    # 2. Swap the Material icon name (the ligature text inside the icon span)
    if icon:
        block = _re.sub(
            r'(?is)(data-loree-role="iconTag"[^>]*>\s*(?:<span\b[^>]*>\s*)?)'
            r'[A-Za-z_][A-Za-z0-9_ ]*',
            lambda m: m.group(1) + icon, block, count=1)

    # 3. Drop the template's own copy from under the heading. Body content is
    #    written after the banner so it inherits the destination region's
    #    styling instead of the banner's white-on-red type.
    block = _re.sub(r"(?is)<p\b[^>]*>.*?</p>", "", block)
    block = _strip_leaf_lists(block)

    row_open = row_open or (
        '<div class="loree-iframe-content-row row" '
        f'style="padding: 20px 40px; background-color: {NEW_HEADING_BG};" '
        'data-loree-role="row">'
    )
    return (
        f'{row_open}\n'
        '<div class="col-12 loree-iframe-content-column" data-loree-role="column">\n'
        f'{block}\n'
        '</div>\n'
        '</div>'
    )
# ── End banner heading clone ─────────────────────────────────────────────────


def _template_labels(html: str):
    """Every heading, short bold line, and [ ... ] placeholder in the template."""
    out = []
    for m in _re.finditer(r"(?is)<(h[1-6]|strong|b|th)\b[^>]*>(.*?)</\1>", html or ""):
        txt = _re.sub(r"<[^>]+>", " ", m.group(2))
        txt = _htmlesc.unescape(txt).replace("\u00a0", " ")
        txt = _re.sub(r"\s+", " ", txt).strip().rstrip(":")
        if txt and len(txt.split()) <= 8:
            out.append(txt)
    out.extend(_bracket_placeholders(html or ""))
    seen, uniq = set(), []
    for t in out:
        k = _norm_label(t)
        if k and k not in seen:
            seen.add(k)
            uniq.append(t)
    return uniq


def _template_has_label(html: str, name: str) -> bool:
    """True when the template has a heading or [ ... ] spot for this header."""
    target = _norm_label(name)
    if not target or not html:
        return False
    for cand in _template_labels(html):
        c = _norm_label(cand)
        if c == target:
            return True
        if len(c) > 3 and len(target) > 3 and (target in c or c in target):
            return True
    return False


_PROMPT_WORDS = {"prompt", "scenario", "overview", "background", "case"}
_INSTRUCTION_WORDS = {"instruction", "instructions", "direction", "directions",
                      "question", "questions", "task", "tasks", "step", "steps"}
_OBJECTIVE_WORDS = {"objective", "objectives", "purpose", "goal", "goals"}


def classify_header(raw: str, template_html: str = ""):
    """
    Work out where a typed header sends its content.

    Returns one of:
      ("title", "")            the header names the discussion title
      ("template", heading)    the template has a heading of that name
      ("prompt" | "instructions" | "objective", "")   a known section name
      ("new", "")              nothing matches, so the header becomes its own
                               sub-heading inside the instructions region
    """
    key = _norm_label(raw)
    if not key:
        return ("instructions", "")
    if key in _TITLE_LABELS:
        return ("title", "")
    if key in _OVERVIEW_LABELS:
        return ("prompt", "")
    if key in _INSTRUCTION_LABELS:
        return ("instructions", "")
    if key in _OBJECTIVE_LABELS:
        return ("objective", "")
    # A heading the template actually has wins over any guesswork
    if template_html:
        for cand in _template_labels(template_html):
            c = _norm_label(cand)
            if c == key or (len(c) > 3 and len(key) > 3
                            and (key in c or c in key)):
                return ("template", cand)
    # "This Week's Prompt" is still the prompt; "Discussion Questions" is still
    # the instructions. Match on the meaningful word inside the header.
    tokens = set(_re.findall(r"[a-z]+", key))
    if tokens & _PROMPT_WORDS:
        return ("prompt", "")
    if tokens & _INSTRUCTION_WORDS:
        return ("instructions", "")
    if tokens & _OBJECTIVE_WORDS:
        return ("objective", "")
    return ("new", "")


def sections_from_pairs(pairs, title: str = "", template_html: str = "",
                        fold_unmatched: bool = True):
    """
    Build (title, body_html, sections) from a list of (header, content) pairs.

    Output matches parse_discussion_docx exactly, so everything downstream —
    the preview, the merge, the created discussion — works the same whether the
    content came from a Word doc or was typed and pasted here.
    """
    resolved_title = (title or "").strip()
    prompt, objective, instructions, extras, unmatched = [], [], [], [], []
    banner = None

    # If one of the sections is named after the banner's own heading, that
    # section owns the banner and no unmatched header may take it.
    banner_taken = False
    _banner_key = _norm_label(banner_heading_text(template_html))
    if _banner_key:
        for _h, _c in pairs:
            _route, _match = classify_header((_h or "").strip().rstrip(":"),
                                             template_html)
            if _route == "template" and _norm_label(_match) == _banner_key:
                banner_taken = True
                break

    order = {}
    labels = {}
    # Two sections can route to the same region ("Instructions" for an
    # unlabelled opener and again for "Discussion Questions"). Joining them into
    # one blob loses the second one's header and makes both inherit the first
    # one's place in the doc, so each section is also kept on its own here.
    chunks = []

    def _mark(key, idx, shown=""):
        if key not in order:
            order[key] = idx
            if shown:
                labels[key] = shown

    def _chunk(route, idx, shown, html):
        chunks.append({"route": route, "order": idx,
                       "label": shown or "", "html": html})

    for _pi, (header, content) in enumerate(pairs):
        shown = (header or "").strip()
        raw = (header or "").strip().rstrip(":")
        html = _text_to_html(content)
        route, _match = classify_header(raw, template_html)

        if route == "title":
            if not resolved_title:
                plain = _html_to_lines(html)
                resolved_title = plain[0].lstrip("• ").strip() if plain else ""
            continue
        if not html:
            continue

        if route == "prompt":
            prompt.append(html)
            _mark("prompt", _pi, shown)
            _chunk("prompt", _pi, shown, html)
        elif route == "objective":
            objective.append(html)
            _mark("objective", _pi, shown)
            _chunk("objective", _pi, shown, html)
        elif route == "instructions":
            instructions.append(html)
            _mark("instructions", _pi, shown)
            _chunk("instructions", _pi, shown, html)
        elif route == "template":
            extras.append({"label": raw, "keyword": _norm_label(raw),
                           "bracket": raw.startswith("[") and raw.endswith("]"),
                           "order": _pi, "display": shown, "html": html})
        elif fold_unmatched:
            # Nothing in the template is named this. The first such header takes
            # over the template's banner: its text becomes the banner heading and
            # its content replaces the placeholder line underneath, so it comes
            # out white-on-red like the rest of the banner.
            if banner is None and not banner_taken:
                banner = {"label": raw.strip("[]").strip(), "html": html,
                          "order": _pi}
                _mark("_banner", _pi)
            else:
                # Later ones cannot share the one banner, so they get a copy of
                # it built inline instead of a bare, unstyled heading.
                unmatched.append(raw)
                instructions.append(
                    f"{banner_heading_html(template_html, raw)}\n{html}"
                )
                _mark("instructions", _pi, shown)
                _chunk("instructions", _pi, shown,
                       f"{banner_heading_html(template_html, raw)}\n{html}")
        else:
            extras.append({"label": raw, "keyword": _norm_label(raw),
                           "bracket": raw.startswith("[") and raw.endswith("]"),
                           "order": _pi, "display": shown, "html": html})

    week = None
    m = _re.search(r"week\s*#?\s*(\d+)", resolved_title, _re.I)
    if m:
        week = int(m.group(1))

    sections = {
        "_week": week,
        "_topic": resolved_title,
        "_order": order,
        "_labels": labels,
        "_chunks": chunks,
        "_regions": extras,
        "_typed": True,
        "_unmatched": unmatched,
        "_banner": banner,
        "prompt": "\n".join(p for p in prompt if p),
        "objective": "\n".join(o for o in objective if o),
        "instructions": "\n".join(i for i in instructions if i),
    }

    body_html = "\n".join(
        chunk for chunk in
        ([banner["html"] if banner else "",
          sections["prompt"], sections["objective"], sections["instructions"]]
         + [e["html"] for e in extras])
        if chunk
    )
    return resolved_title or "Untitled Discussion", body_html, sections
# ── End typed headers + pasted content ───────────────────────────────────────


def _html_to_lines(html: str):
    """Flatten a chunk of HTML into plain lines that are easy to copy out."""
    out = []
    for m in _re.finditer(r"(?is)<(p|li)\b[^>]*>(.*?)</\1>", html or ""):
        text = _re.sub(r"<[^>]+>", "", m.group(2))
        text = _htmlesc.unescape(text)
        text = _re.sub(r"\s+", " ", text).strip()
        if text:
            out.append(("• " if m.group(1).lower() == "li" else "") + text)
    if not out:
        text = _re.sub(r"\s+", " ", _re.sub(r"<[^>]+>", " ", html or "")).strip()
        if text:
            out.append(text)
    return out


def normalize_template_html(raw: str) -> str:
    """
    Clean up HTML pasted straight out of the Canvas HTML editor (or a saved
    .html file). Strips a full-document wrapper if present and keeps any
    <style> blocks that were sitting in the <head>.
    """
    html = (raw or "").strip().lstrip("\ufeff")
    if not html:
        return ""

    styles = _re.findall(r"<style\b[^>]*>.*?</style>", html, _re.I | _re.S)
    m = _re.search(r"<body\b[^>]*>(.*?)</body>", html, _re.I | _re.S)
    if m:
        html = m.group(1)
        for s in styles:
            if s not in html:
                html = s + "\n" + html
    return html.strip()


def template_headings(html: str):
    """Every heading in the template, plain text — used by the Check button."""
    out = []
    for m in _re.finditer(r"(?is)<h[1-6]\b[^>]*>(.*?)</h[1-6]>", html or ""):
        txt = _re.sub(r"<[^>]+>", " ", m.group(1))
        txt = _htmlesc.unescape(txt).replace("\u00a0", " ")
        txt = _re.sub(r"\s+", " ", txt).strip()
        if txt:
            out.append(txt)
    return out


# The section headings a template is allowed to have. These double as the
# boundaries between regions and as the headings the banner rename must skip.
TEMPLATE_SECTION_LABELS = [
    "topic introduction", "submissions", "this week's prompt",
    "discussion prompt", "prompt", "scenario", "overview", "background",
    "summary", "discussion instructions", "instructions", "directions",
    "task", "requirements", "questions", "discussion questions",
    "initial post", "response requirements", "step 1", "step 2",
    "helpful tips", "helpful tip", "tips", "resources", "rubric",
    "evaluation", "grading", "objective", "objectives",
]


def parse_docx_bytes(data: bytes, filename: str):
    """parse_discussion_docx for an uploaded file instead of a path on disk."""
    import tempfile
    import os

    suffix = Path(filename or "upload.docx").suffix or ".docx"
    fd, tmp = tempfile.mkstemp(suffix=suffix)
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        title, body_html, sections = parse_discussion_docx(tmp)
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass
    # The title normally comes from the doc's own header lines; fall back to
    # the uploaded name so an entry is never left untitled.
    if not title:
        title = _title_from_filename(filename or "Discussion")
    return title, body_html, sections


# ── titles ───────────────────────────────────────────────────────────────────
def apply_week_prefix(title: str, week) -> str:
    """"Medicare Matters" + week 3 → "Week 3 Discussion: Medicare Matters"."""
    t = (title or "").strip()
    if week and not _re.match(r"^\s*week\s*\d", t, _re.I):
        return f"Week {week} Discussion: {t}"
    return t


def title_for_page(title: str) -> str:
    """
    The heading that goes inside the page, taken off the Canvas title.

    The Canvas discussion title carries the week number; the heading inside the
    page does not. Deriving one from the other means an edit in the titles box
    shows up in both places.
    """
    page = _re.sub(r"^\s*(?:week\s*\d+\s*)?discussion\s*[:\-\u2010-\u2015]\s*", "",
                   title or "", flags=_re.I).strip()
    page = _re.sub(r"^\s*week\s*\d+\s*[:\-\u2010-\u2015]\s*", "", page,
                   flags=_re.I).strip()
    return page or (title or "").strip()


# Public name for the flattener; merge.py logs with it.
html_to_lines = _html_to_lines
